#!/usr/bin/env python3
"""
SkillGraph Stage 2 — Test Script

Loads Stage 1's output (resume statements + job description), runs
explicit skill extraction via the curated dictionary, and prints the
structured results.
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

from docx import Document

from skillgraph.parser import extract_resume_text, segment_resume, clean_jd_text
from skillgraph.extractor import (
    extract_resume_skills,
    extract_jd_skills,
    load_skill_dictionary,
)

# ── Reuse Stage 1 sample data ───────────────────────────────────────

SAMPLE_RESUME_LINES = [
    ("Priya Sharma", True),
    ("priya.sharma@email.com | linkedin.com/in/priyasharma | github.com/priyaS", False),
    ("", False),
    ("SUMMARY", True),
    (
        "Full-stack engineer with 4+ years of experience building scalable "
        "web applications and data pipelines. Passionate about clean code, "
        "test-driven development, and mentoring junior developers.",
        False,
    ),
    ("", False),
    ("EXPERIENCE", True),
    ("Software Engineer II — Acme Corp, Bangalore (Jan 2022 – Present)", True),
    (
        "• Designed and implemented a microservices architecture using "
        "Python/FastAPI and PostgreSQL, reducing API latency by 35%.",
        False,
    ),
    (
        "• Built a real-time data pipeline with Apache Kafka and Spark "
        "Structured Streaming to process 2M+ events/day.",
        False,
    ),
    (
        "• Led a team of 3 engineers to deliver an internal analytics "
        "dashboard in React and D3.js, adopted by 150+ users.",
        False,
    ),
    (
        "• Wrote comprehensive unit and integration tests using pytest, "
        "achieving 92% code coverage.",
        False,
    ),
    ("", False),
    ("Software Engineer — StartupXYZ, Hyderabad (Jun 2020 – Dec 2021)", True),
    (
        "• Developed RESTful APIs in Django REST Framework serving 50k "
        "daily active users.",
        False,
    ),
    (
        "• Automated CI/CD pipelines with GitHub Actions and Docker, "
        "cutting deployment time from 45 min to 8 min.",
        False,
    ),
    (
        "• Integrated Elasticsearch for full-text search across 1.2M "
        "product listings, improving search speed by 60%.",
        False,
    ),
    ("", False),
    ("PROJECTS", True),
    ("SkillGraph (Personal Project)", True),
    (
        "• Building an NLP pipeline to infer implicit skills from resumes "
        "using spaCy, sentence transformers, and knowledge graphs.",
        False,
    ),
    ("", False),
    ("Open-source Contributor — Apache Airflow", True),
    (
        "• Fixed 12 bugs related to DAG scheduling and task retries; "
        "reviewed 30+ community pull requests.",
        False,
    ),
    ("", False),
    ("EDUCATION", True),
    ("B.Tech in Computer Science — IIT Hyderabad (2016 – 2020)", False),
    ("", False),
    ("SKILLS", True),
    (
        "Python, JavaScript, TypeScript, SQL, FastAPI, Django, React, "
        "Docker, Kubernetes, Kafka, Spark, PostgreSQL, Redis, Git",
        False,
    ),
    ("", False),
    ("CERTIFICATIONS", True),
    ("AWS Certified Solutions Architect – Associate (2023)", False),
    ("", False),
    ("---", False),
]

SAMPLE_JD = """\
Senior Backend Engineer — FinTech Startup (Remote)

About the role
We're looking for a Senior Backend Engineer to design,  build, and
maintain high-throughput financial systems.   You'll work closely with
the product team to ship features that handle millions of transactions
per day.

Responsibilities
- Architect and implement scalable backend services in Python
- Design event-driven systems using  message queues (Kafka, RabbitMQ)
- Own the full lifecycle of features from design to production
- Mentor junior engineers and conduct code reviews
- Collaborate with data engineering to build reliable ETL pipelines

Requirements
- 4+ years of backend engineering experience
- Strong proficiency in Python (FastAPI or Django)
- Experience with relational databases (PostgreSQL, MySQL) and caching (Redis)
- Familiarity with containerisation (Docker, Kubernetes)
- Solid understanding of distributed systems and microservices
- Excellent communication skills and ability to work in a remote team

Nice to Have
- Experience with real-time data streaming (Kafka, Spark)
- Contributions to open-source projects
- AWS or GCP certifications
"""


def create_sample_docx(path: Path) -> None:
    """Write a realistic DOCX resume from structured data."""
    doc = Document()
    for text, bold in SAMPLE_RESUME_LINES:
        para = doc.add_paragraph()
        run = para.add_run(text)
        run.bold = bold
    doc.save(str(path))


# ── Main ─────────────────────────────────────────────────────────────

def main() -> None:
    sample_dir = Path(__file__).parent / "samples"
    sample_dir.mkdir(exist_ok=True)
    resume_path = sample_dir / "sample_resume.docx"

    # ── Stage 1: Parsing & Segmentation ──────────────────────────────
    create_sample_docx(resume_path)
    resume_text = extract_resume_text(resume_path)
    statements = segment_resume(resume_text)
    jd_clean = clean_jd_text(SAMPLE_JD)

    # ── Stage 2: Explicit Skill Extraction ───────────────────────────
    skill_dict = load_skill_dictionary()

    resume_skills = extract_resume_skills(statements, skill_dict)
    jd_skills = extract_jd_skills(jd_clean, skill_dict)

    output = {
        "resume_skills": resume_skills,
        "jd_required_skills": jd_skills,
    }

    # ── Human-readable report ────────────────────────────────────────
    print("=" * 64)
    print("  STAGE 2 OUTPUT — Explicit Skill Extraction")
    print("=" * 64)

    print(f"\n🎯  Resume Skills  ({len(resume_skills)} matches)")
    print("-" * 64)
    for entry in resume_skills:
        print(f'  • {entry["skill"]:<25s}  ← "{entry["source_statement"][:60]}…"')

    # Deduplicated skill set
    unique_resume = sorted({e["skill"] for e in resume_skills})
    print(f"\n📊  Unique Resume Skills ({len(unique_resume)}): {', '.join(unique_resume)}")

    print(f"\n📋  JD Required Skills  ({len(jd_skills)} skills)")
    print("-" * 64)
    for skill in jd_skills:
        print(f"  • {skill}")

    # ── JSON output ──────────────────────────────────────────────────
    print("\n" + "=" * 64)
    print("  JSON OUTPUT")
    print("=" * 64)
    print(json.dumps(output, indent=2, ensure_ascii=False))

    return 0


if __name__ == "__main__":
    sys.exit(main() or 0)
