#!/usr/bin/env python3
"""
SkillGraph Stage 3 — Test Script

Runs the full Stage 1 → Stage 2 → Stage 3 chain on the sample resume
and job description, then prints the inference results grouped by
classification.
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

from docx import Document

from skillgraph.parser import extract_resume_text, segment_resume, clean_jd_text
from skillgraph.extractor import (
    extract_resume_skills,
    extract_jd_skills,
    load_skill_dictionary,
)
from skillgraph.inference import (
    build_graph,
    load_skill_graph,
    load_sbert_model,
    run_inference,
)

# ── Sample data (same as Stage 1/2 tests) ───────────────────────────

SAMPLE_RESUME_LINES = [
    ("Priya Sharma", True),
    ("priya.sharma@email.com | linkedin.com/in/priyasharma | github.com/priyaS", False),
    ("", False),
    ("SUMMARY", True),
    (
        "Full-stack engineer with 4+ years of experience building scalable "
        "web applications and data pipelines. Passionate about clean code, "
        "test-driven development, and mentoring junior developers.",
        False,
    ),
    ("", False),
    ("EXPERIENCE", True),
    ("Software Engineer II — Acme Corp, Bangalore (Jan 2022 – Present)", True),
    (
        "• Designed and implemented a microservices architecture using "
        "Python/FastAPI and PostgreSQL, reducing API latency by 35%.",
        False,
    ),
    (
        "• Built a real-time data pipeline with Apache Kafka and Spark "
        "Structured Streaming to process 2M+ events/day.",
        False,
    ),
    (
        "• Led a team of 3 engineers to deliver an internal analytics "
        "dashboard in React and D3.js, adopted by 150+ users.",
        False,
    ),
    (
        "• Wrote comprehensive unit and integration tests using pytest, "
        "achieving 92% code coverage.",
        False,
    ),
    ("", False),
    ("Software Engineer — StartupXYZ, Hyderabad (Jun 2020 – Dec 2021)", True),
    (
        "• Developed RESTful APIs in Django REST Framework serving 50k "
        "daily active users.",
        False,
    ),
    (
        "• Automated CI/CD pipelines with GitHub Actions and Docker, "
        "cutting deployment time from 45 min to 8 min.",
        False,
    ),
    (
        "• Integrated Elasticsearch for full-text search across 1.2M "
        "product listings, improving search speed by 60%.",
        False,
    ),
    ("", False),
    ("PROJECTS", True),
    ("SkillGraph (Personal Project)", True),
    (
        "• Building an NLP pipeline to infer implicit skills from resumes "
        "using spaCy, sentence transformers, and knowledge graphs.",
        False,
    ),
    ("", False),
    ("Open-source Contributor — Apache Airflow", True),
    (
        "• Fixed 12 bugs related to DAG scheduling and task retries; "
        "reviewed 30+ community pull requests.",
        False,
    ),
    ("", False),
    ("EDUCATION", True),
    ("B.Tech in Computer Science — IIT Hyderabad (2016 – 2020)", False),
    ("", False),
    ("SKILLS", True),
    (
        "Python, JavaScript, TypeScript, SQL, FastAPI, Django, React, "
        "Docker, Kubernetes, Kafka, Spark, PostgreSQL, Redis, Git",
        False,
    ),
    ("", False),
    ("CERTIFICATIONS", True),
    ("AWS Certified Solutions Architect – Associate (2023)", False),
    ("", False),
    ("---", False),
]

SAMPLE_JD = """\
Senior Backend Engineer — FinTech Startup (Remote)

About the role
We're looking for a Senior Backend Engineer to design,  build, and
maintain high-throughput financial systems.   You'll work closely with
the product team to ship features that handle millions of transactions
per day.

Responsibilities
- Architect and implement scalable backend services in Python
- Design event-driven systems using  message queues (Kafka, RabbitMQ)
- Own the full lifecycle of features from design to production
- Mentor junior engineers and conduct code reviews
- Collaborate with data engineering to build reliable ETL pipelines

Requirements
- 4+ years of backend engineering experience
- Strong proficiency in Python (FastAPI or Django)
- Experience with relational databases (PostgreSQL, MySQL) and caching (Redis)
- Familiarity with containerisation (Docker, Kubernetes)
- Solid understanding of distributed systems and microservices
- Excellent communication skills and ability to work in a remote team

Nice to Have
- Experience with real-time data streaming (Kafka, Spark)
- Contributions to open-source projects
- AWS or GCP certifications
"""


def create_sample_docx(path: Path) -> None:
    """Write a realistic DOCX resume from structured data."""
    doc = Document()
    for text, bold in SAMPLE_RESUME_LINES:
        para = doc.add_paragraph()
        run = para.add_run(text)
        run.bold = bold
    doc.save(str(path))


# ── Main ─────────────────────────────────────────────────────────────

def main() -> None:
    sample_dir = Path(__file__).parent / "samples"
    sample_dir.mkdir(exist_ok=True)
    resume_path = sample_dir / "sample_resume.docx"

    # ── Stage 1: Parsing & Segmentation ──────────────────────────────
    print("⏳ Stage 1: Parsing & Segmentation...")
    create_sample_docx(resume_path)
    resume_text = extract_resume_text(resume_path)
    statements = segment_resume(resume_text)
    jd_clean = clean_jd_text(SAMPLE_JD)
    print(f"   ✓ {len(statements)} resume statements extracted\n")

    # ── Stage 2: Explicit Skill Extraction ───────────────────────────
    print("⏳ Stage 2: Explicit Skill Extraction...")
    skill_dict = load_skill_dictionary()
    resume_skills = extract_resume_skills(statements, skill_dict)
    jd_skills = extract_jd_skills(jd_clean, skill_dict)

    stage2_output = {
        "resume_skills": resume_skills,
        "jd_required_skills": jd_skills,
    }

    unique_resume = sorted({e["skill"] for e in resume_skills})
    print(f"   ✓ {len(unique_resume)} unique resume skills")
    print(f"   ✓ {len(jd_skills)} JD required skills\n")

    # ── Stage 3: Inference Engine ────────────────────────────────────
    print("⏳ Stage 3: Loading skill graph...")
    edges = load_skill_graph()
    graph = build_graph(edges)
    print(f"   ✓ Graph: {graph.number_of_nodes()} nodes, "
          f"{graph.number_of_edges()} edges")

    print("⏳ Stage 3: Loading Sentence-BERT model (first run downloads ~80 MB)...")
    model = load_sbert_model()
    print("   ✓ Model loaded\n")

    print("⏳ Stage 3: Running inference...\n")
    result = run_inference(stage2_output, graph, model)

    # ── Display results ──────────────────────────────────────────────
    print("=" * 70)
    print("  STAGE 3 OUTPUT — Inference Engine")
    print("=" * 70)

    # Explicitly matched skills
    print(f"\n✅  Explicitly Matched Skills ({len(result['explicitly_matched'])})")
    print("-" * 70)
    for skill in result["explicitly_matched"]:
        print(f"   • {skill}")

    # Group inferences by classification
    groups: dict[str, list] = {
        "HIGH_CONFIDENCE": [],
        "NEEDS_VERIFICATION": [],
        "SKILL_GAP": [],
    }
    for inf in result["inferences"]:
        groups[inf["classification"]].append(inf)

    # HIGH_CONFIDENCE
    print(f"\n🟢  HIGH CONFIDENCE — Inferred Skills "
          f"({len(groups['HIGH_CONFIDENCE'])})")
    print("-" * 70)
    if not groups["HIGH_CONFIDENCE"]:
        print("   (none)")
    for inf in groups["HIGH_CONFIDENCE"]:
        _print_inference(inf)

    # NEEDS_VERIFICATION
    print(f"\n🟡  NEEDS VERIFICATION — Candidate Should Confirm "
          f"({len(groups['NEEDS_VERIFICATION'])})")
    print("-" * 70)
    if not groups["NEEDS_VERIFICATION"]:
        print("   (none)")
    for inf in groups["NEEDS_VERIFICATION"]:
        _print_inference(inf)

    # SKILL_GAP
    print(f"\n🔴  SKILL GAP — Not Found or No Evidence "
          f"({len(groups['SKILL_GAP'])})")
    print("-" * 70)
    if not groups["SKILL_GAP"]:
        print("   (none)")
    for inf in groups["SKILL_GAP"]:
        _print_inference(inf)

    # ── JSON output ──────────────────────────────────────────────────
    print("\n" + "=" * 70)
    print("  JSON OUTPUT")
    print("=" * 70)
    print(json.dumps(result, indent=2, ensure_ascii=False))

    return 0


def _print_inference(inf: dict) -> None:
    """Pretty-print a single inference object."""
    print(f"\n   📌 {inf['missing_skill']}  "
          f"[confidence: {inf['confidence']}%]")
    if inf["matched_resume_skill"]:
        print(f"      Linked to: {inf['matched_resume_skill']}")
        stmt = inf["matched_statement"]
        if len(stmt) > 80:
            stmt = stmt[:77] + "..."
        print(f"      Statement: \"{stmt}\"")
        print(f"      Edge weight: {inf['edge_weight']}  |  "
              f"Semantic sim: {inf['semantic_similarity']}")
    print(f"      💬 {inf['reasoning']}")


if __name__ == "__main__":
    sys.exit(main() or 0)
