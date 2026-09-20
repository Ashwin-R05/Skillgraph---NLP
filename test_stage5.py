#!/usr/bin/env python3
"""
SkillGraph — Full End-to-End Pipeline Test (Stages 1 through 5)

Executes all 5 stages of the SkillGraph pipeline sequentially:
  - Stage 1: Parsing & Segmentation (DOCX parsing, statement segmentation, JD cleaning)
  - Stage 2: Explicit Skill Extraction (NER + curated skill dictionary)
  - Stage 3: Inference Engine (Graph traversal + Sentence-BERT semantic similarity scoring)
  - Stage 4: Clarifying Question Generation & Verification (Template generation & candidate response simulation)
  - Stage 5: Report Generation & Resume Rewrite Suggestions (Fit breakdown, score compilation, constrained rewriting)
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

from docx import Document

# Stage 1 imports
from skillgraph.parser import extract_resume_text, segment_resume, clean_jd_text

# Stage 2 imports
from skillgraph.extractor import (
    extract_resume_skills,
    extract_jd_skills,
    load_skill_dictionary,
)

# Stage 3 imports
from skillgraph.inference import (
    build_graph,
    load_skill_graph,
    load_sbert_model,
    run_inference,
)

# Stage 4 imports
from skillgraph.verifier import (
    load_question_templates,
    generate_all_questions,
    apply_response,
    initialize_verification_status,
)

# Stage 5 imports
from skillgraph.reporter import (
    compile_report,
    generate_all_rewrites,
    build_final_output,
)


# ── Sample Data (Consistent across all stages) ───────────────────────

SAMPLE_RESUME_LINES = [
    ("Priya Sharma", True),
    ("priya.sharma@email.com | linkedin.com/in/priyasharma | github.com/priyaS", False),
    ("", False),
    ("SUMMARY", True),
    (
        "Full-stack engineer with 4+ years of experience building scalable "
        "web applications and data pipelines. Passionate about clean code, "
        "test-driven development, and mentoring junior developers.",
        False,
    ),
    ("", False),
    ("EXPERIENCE", True),
    ("Software Engineer II — Acme Corp, Bangalore (Jan 2022 – Present)", True),
    (
        "• Designed and implemented a microservices architecture using "
        "Python/FastAPI and PostgreSQL, reducing API latency by 35%.",
        False,
    ),
    (
        "• Built a real-time data pipeline with Apache Kafka and Spark "
        "Structured Streaming to process 2M+ events/day.",
        False,
    ),
    (
        "• Led a team of 3 engineers to deliver an internal analytics "
        "dashboard in React and D3.js, adopted by 150+ users.",
        False,
    ),
    (
        "• Wrote comprehensive unit and integration tests using pytest, "
        "achieving 92% code coverage.",
        False,
    ),
    ("", False),
    ("Software Engineer — StartupXYZ, Hyderabad (Jun 2020 – Dec 2021)", True),
    (
        "• Developed RESTful APIs in Django REST Framework serving 50k "
        "daily active users.",
        False,
    ),
    (
        "• Automated CI/CD pipelines with GitHub Actions and Docker, "
        "cutting deployment time from 45 min to 8 min.",
        False,
    ),
    (
        "• Integrated Elasticsearch for full-text search across 1.2M "
        "product listings, improving search speed by 60%.",
        False,
    ),
    ("", False),
    ("PROJECTS", True),
    ("SkillGraph (Personal Project)", True),
    (
        "• Building an NLP pipeline to infer implicit skills from resumes "
        "using spaCy, sentence transformers, and knowledge graphs.",
        False,
    ),
    ("", False),
    ("Open-source Contributor — Apache Airflow", True),
    (
        "• Fixed 12 bugs related to DAG scheduling and task retries; "
        "reviewed 30+ community pull requests.",
        False,
    ),
    ("", False),
    ("EDUCATION", True),
    ("B.Tech in Computer Science — IIT Hyderabad (2016 – 2020)", False),
    ("", False),
    ("SKILLS", True),
    (
        "Python, JavaScript, TypeScript, SQL, FastAPI, Django, React, "
        "Docker, Kubernetes, Kafka, Spark, PostgreSQL, Redis, Git",
        False,
    ),
    ("", False),
    ("CERTIFICATIONS", True),
    ("AWS Certified Solutions Architect – Associate (2023)", False),
    ("", False),
    ("---", False),
]

SAMPLE_JD = """\
Senior Backend Engineer — FinTech Startup (Remote)

About the role
We're looking for a Senior Backend Engineer to design,  build, and
maintain high-throughput financial systems.   You'll work closely with
the product team to ship features that handle millions of transactions
per day.

Responsibilities
- Architect and implement scalable backend services in Python
- Design event-driven systems using  message queues (Kafka, RabbitMQ)
- Own the full lifecycle of features from design to production
- Mentor junior engineers and conduct code reviews
- Collaborate with data engineering to build reliable ETL pipelines

Requirements
- 4+ years of backend engineering experience
- Strong proficiency in Python (FastAPI or Django)
- Experience with relational databases (PostgreSQL, MySQL) and caching (Redis)
- Familiarity with containerisation (Docker, Kubernetes)
- Solid understanding of distributed systems and microservices
- Excellent communication skills and ability to work in a remote team

Nice to Have
- Experience with real-time data streaming (Kafka, Spark)
- Contributions to open-source projects
- AWS or GCP certifications
"""


def create_sample_docx(path: Path) -> None:
    doc = Document()
    for text, bold in SAMPLE_RESUME_LINES:
        para = doc.add_paragraph()
        run = para.add_run(text)
        run.bold = bold
    doc.save(str(path))


def main() -> None:
    sample_dir = Path(__file__).parent / "samples"
    sample_dir.mkdir(exist_ok=True)
    resume_path = sample_dir / "sample_resume.docx"

    print("=" * 76)
    print("  SKILLGRAPH PIPELINE — FULL END-TO-END DEMONSTRATION (STAGES 1 TO 5)")
    print("=" * 76)

    # ── Stage 1: Parsing & Segmentation ──────────────────────────────
    print("\n[1/5] Stage 1: Parsing & Segmentation...")
    create_sample_docx(resume_path)
    resume_text = extract_resume_text(resume_path)
    statements = segment_resume(resume_text)
    jd_clean = clean_jd_text(SAMPLE_JD)
    print(f"      ✓ Extracted {len(statements)} resume statements.")
    print(f"      ✓ Cleaned job description ({len(jd_clean)} characters).")

    # ── Stage 2: Explicit Skill Extraction ───────────────────────────
    print("\n[2/5] Stage 2: Explicit Skill Extraction...")
    skill_dict = load_skill_dictionary()
    resume_skills = extract_resume_skills(statements, skill_dict)
    jd_skills = extract_jd_skills(jd_clean, skill_dict)

    stage2_output = {
        "resume_skills": resume_skills,
        "jd_required_skills": jd_skills,
    }
    unique_resume = sorted({e["skill"] for e in resume_skills})
    print(f"      ✓ Identified {len(unique_resume)} unique skills in resume.")
    print(f"      ✓ Identified {len(jd_skills)} required skills in job description.")

    # ── Stage 3: Inference Engine ────────────────────────────────────
    print("\n[3/5] Stage 3: Inference Engine...")
    edges = load_skill_graph()
    graph = build_graph(edges)
    model = load_sbert_model()
    stage3_output = run_inference(stage2_output, graph, model)
    inferences = stage3_output["inferences"]
    print(f"      ✓ Evaluated {len(inferences)} missing skills against skill graph & SBERT.")

    # ── Stage 4: Clarifying Question Generation & Verification ──────
    print("\n[4/5] Stage 4: Clarifying Question Generation & Verification...")
    inferences = initialize_verification_status(inferences)
    templates = load_question_templates()
    questions = generate_all_questions(inferences, templates)
    print(f"      ✓ Generated {len(questions)} clarifying questions for candidate verification.")

    # Candidate Verification Simulation:
    # 1. Distributed Systems -> "Implemented it myself" (CONFIRMED)
    # 2. ETL -> "Used a library/tool" (PARTIALLY_CONFIRMED)
    # 3. MySQL -> "Not sure" (SKILL_GAP)
    simulated_answers = {
        "Distributed Systems": "Implemented it myself",
        "ETL": "Used a library/tool",
        "MySQL": "Not sure",
    }

    updated_inferences = []
    for inf in inferences:
        skill_name = inf.get("missing_skill")
        if skill_name in simulated_answers:
            ans = simulated_answers[skill_name]
            updated = apply_response(inf, ans)
            updated_inferences.append(updated)
            print(f"      • Answered '{skill_name}': '{ans}' -> {updated['classification']}")
        else:
            updated_inferences.append(inf)

    stage4_output = {
        "inferences": updated_inferences,
        "explicitly_matched": stage3_output["explicitly_matched"],
    }

    # ── Stage 5: Report Generation & Resume Rewrite Suggestions ──────
    print("\n[5/5] Stage 5: Report Generation & Resume Rewrite Suggestions...")
    report = compile_report(stage4_output)
    rewrites = generate_all_rewrites(report)
    final_output = build_final_output(report, rewrites)

    # ── Print Human-Readable Summary ─────────────────────────────────
    print("\n" + "=" * 76)
    print("  FINAL CANDIDATE FIT & REWRITE REPORT")
    print("=" * 76)

    fit_summary = final_output["fit_summary"]
    print(f"\n📊 Overall Match Summary:")
    print(f"   • Total Required Skills : {fit_summary['total_required_skills']}")
    print(f"   • Matched Skills Count  : {fit_summary['matched_count']}")
    print(f"   • Candidate Fit Score   : {fit_summary['fit_percentage']}%")

    breakdown = final_output["skills_breakdown"]
    print(f"\n✅ Explicitly Present ({len(breakdown['explicitly_present'])}):")
    for s in breakdown["explicitly_present"]:
        print(f"   [EXPLICIT]            {s}")

    print(f"\n🟢 Confirmed by Verification ({len(breakdown['confirmed_by_inference'])}):")
    for inf in breakdown["confirmed_by_inference"]:
        print(f"   [CONFIRMED]           {inf['missing_skill']} (Confidence: {inf['confidence']}%, Reason: {inf['reasoning'][:65]}...)")

    print(f"\n🟡 Partially Confirmed ({len(breakdown['partially_confirmed'])}):")
    for inf in breakdown["partially_confirmed"]:
        print(f"   [PARTIAL]             {inf['missing_skill']} (Confidence: {inf['confidence']}%, Reason: {inf['reasoning'][:65]}...)")

    print(f"\n🔵 High Confidence (Unverified) ({len(breakdown['high_confidence_unverified'])}):")
    if not breakdown["high_confidence_unverified"]:
        print("   (none)")
    for inf in breakdown["high_confidence_unverified"]:
        print(f"   [HIGH_CONF_UNVERIFIED] {inf['missing_skill']}")

    print(f"\n🔴 Genuine Skill Gaps ({len(breakdown['genuine_gaps'])}):")
    for inf in breakdown["genuine_gaps"]:
        print(f"   [SKILL_GAP]           {inf['missing_skill']}")

    print(f"\n✍️ Suggested Resume Bullet Rewrites ({len(final_output['suggested_rewrites'])}):")
    print("-" * 76)
    for i, rw in enumerate(final_output["suggested_rewrites"], 1):
        print(f"\n[{i}] Target Skill: {rw['missing_skill']}")
        print(f"    Original : \"{rw['original_statement']}\"")
        print(f"    Suggested: \"{rw['suggested_rewrite']}\"")

    # ── Full JSON Dump ───────────────────────────────────────────────
    print("\n" + "=" * 76)
    print("  FINAL JSON OUTPUT OBJECT")
    print("=" * 76)
    print(json.dumps(final_output, indent=2, ensure_ascii=False))

    return 0


if __name__ == "__main__":
    sys.exit(main() or 0)
